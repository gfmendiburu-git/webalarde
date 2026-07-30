#!/usr/bin/env python3
"""Synchronize cantineras JSON data with the working Excel file."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

import openpyxl


DEFAULT_EXCEL = Path("/home/gaizka/Descargas/CANTINERAS_actualizado.xlsx")
DEFAULT_JSON = Path("data/cantineras.json")


def normalized(value: object) -> str:
    text = str(value or "").strip().lower()
    text = "".join(
        char for char in unicodedata.normalize("NFD", text)
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", text)


def truthy(value: object) -> bool:
    return normalized(value) in {"si", "sí", "yes", "true", "1", "x"}


def clean(value: object) -> str:
    return str(value or "").strip()


def entry_key(year: object, company: object) -> tuple[int, str]:
    return int(year), normalized(company)


def read_excel(path: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Cantineras"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, object]] = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in raw):
            continue
        rows.append(dict(zip(headers, raw)))
    return rows


def sync_data(data: dict[str, object], rows: list[dict[str, object]]) -> tuple[dict[str, object], dict[str, int]]:
    entries = data.get("entries", [])
    if not isinstance(entries, list):
        raise ValueError("data/cantineras.json no contiene una lista de entries")

    existing = {entry_key(entry["year"], entry["company"]): entry for entry in entries}
    synced: list[dict[str, object]] = []
    added = 0
    changed = 0

    for row in rows:
        key = entry_key(row["año"], row["compañía"])
        base = dict(existing.get(key, {}))
        if not base:
            added += 1
            base = {
                "year": int(row["año"]),
                "company": clean(row["compañía"]),
                "name": clean(row["nombre"]),
                "source_url": "",
                "source_title": "",
                "mentions": 1,
                "needs_review": True,
                "alternatives": [],
                "confirmed": False,
                "pending_confirmation": True,
            }

        before = json.dumps(base, ensure_ascii=False, sort_keys=True)
        pending = truthy(row.get("pendiente_confirmacion"))
        base["year"] = int(row["año"])
        base["company"] = clean(row["compañía"])
        base["name"] = clean(row["nombre"])
        base["confirmed"] = truthy(row.get("confirmada"))
        base["pending_confirmation"] = pending
        base["needs_review"] = True if pending else truthy(row.get("revisar"))
        if row.get("menciones") not in (None, ""):
            base["mentions"] = int(row["menciones"])
        elif not base.get("mentions"):
            base["mentions"] = 1
        base["source_title"] = clean(row.get("fuente"))
        base["source_url"] = clean(row.get("url_fuente"))
        base.setdefault("alternatives", [])

        after = json.dumps(base, ensure_ascii=False, sort_keys=True)
        if before != after and key in existing:
            changed += 1
        synced.append(base)

    synced.sort(key=lambda item: (int(item["year"]), normalized(item["company"]), normalized(item["name"])))
    data["entries"] = synced
    stats = data.setdefault("stats", {})
    if isinstance(stats, dict):
        stats["entries"] = len(synced)
        stats["total_entries"] = stats["entries"]
        stats["confirmed"] = sum(1 for entry in synced if entry.get("confirmed"))
        stats["pending_confirmation"] = sum(1 for entry in synced if entry.get("pending_confirmation"))
        stats["needs_review"] = sum(1 for entry in synced if entry.get("needs_review"))
        stats["confirmed_entries"] = stats["confirmed"]
        stats["pending_confirmation_entries"] = stats["pending_confirmation"]
        stats["needs_review_entries"] = stats["needs_review"]
    return data, {"added": added, "changed": changed, "entries": len(synced)}


def export_excel(data: dict[str, object], path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Cantineras"
    headers = [
        "año",
        "compañía",
        "nombre",
        "confirmada",
        "pendiente_confirmacion",
        "revisar",
        "menciones",
        "fuente",
        "url_fuente",
        "alternativas",
    ]
    sheet.append(headers)
    for entry in data["entries"]:
        alternatives = entry.get("alternatives") or []
        sheet.append([
            entry.get("year", ""),
            entry.get("company", ""),
            entry.get("name", ""),
            "sí" if entry.get("confirmed") else "no",
            "sí" if entry.get("pending_confirmation") else "no",
            "sí" if entry.get("needs_review") else "no",
            entry.get("mentions", ""),
            entry.get("source_title", ""),
            entry.get("source_url", ""),
            "; ".join(
                f"{item.get('name', '')} ({item.get('source_title', '')})".strip()
                for item in alternatives
            ),
        ])

    summary = workbook.create_sheet("Resumen")
    stats = data.get("stats", {})
    summary.append(["métrica", "valor"])
    if isinstance(stats, dict):
        for key in ("entries", "confirmed", "pending_confirmation", "needs_review"):
            summary.append([key, stats.get(key, "")])

    workbook.save(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--json", type=Path, default=DEFAULT_JSON)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    rows = read_excel(args.excel)
    data = json.loads(args.json.read_text(encoding="utf-8"))
    data, stats = sync_data(data, rows)
    args.json.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    export_excel(data, args.excel)
    print(f"{stats['entries']} registros sincronizados")
    print(f"{stats['added']} altas; {stats['changed']} cambios")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
